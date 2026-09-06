"""Reading and normalisation of the IR/ATR spectra."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from . import descriptions

# Common resampling grid. Every worksheet in the source workbook covers
# approximately 499-4000 cm^-1; this range is shared by all of them.
WAVENUMBER_MIN = 500.0
WAVENUMBER_MAX = 3996.0
N_GRID_POINTS = 3600

# Material family assignment, applied to the free-text description of each
# reference. Order matters: the first matching pattern wins.
FAMILY_PATTERNS: list[tuple[str, str]] = [
    ("TPU", r"\bTPU\b"),
    ("PUR", r"\bPUR\b|Hi-?react\s*PU\b|\bPU\b"),
    ("TR", r"\bTR\b"),
    ("EVA", r"\bEVA\b"),
    ("PVC", r"\bPVC\b"),
    ("RUBBER", r"caucho|LATEX"),
    ("LEATHER", r"cuerolite|Piel"),
]

# Worksheets whose name is not a reference number. 'H4965' is the TR grade
# "H49 65" (reference 47, supplier Ruiz Alejos).
SPECIAL_SHEETS: dict[str, tuple[int, str, str]] = {
    "H4965": (47, "15 planchas TR negro H49 65", "Ruiz Alejos"),
}


@dataclass
class SpectraSet:
    """A collection of IR spectra resampled on a common wavenumber grid."""

    wavenumber: np.ndarray  # (n_points,)
    intensity: np.ndarray  # (n_samples, n_points)
    metadata: pd.DataFrame  # one row per sample

    def __len__(self) -> int:
        return self.intensity.shape[0]


# ---------------------------------------------------------------------------
# Reference list
# ---------------------------------------------------------------------------
def read_references(path: str | Path) -> dict[int, tuple[str, str]]:
    """Parse ``references.txt`` into ``{reference_number: (description, supplier)}``.

    The file is a three-column table (number / description / supplier) exported
    with one field per line.
    """
    text = Path(path).read_text(encoding="utf-8", errors="replace")
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]

    refs: dict[int, tuple[str, str]] = {}
    i = 0
    while i + 2 < len(lines):
        if re.fullmatch(r"\d+", lines[i]):
            refs[int(lines[i])] = (lines[i + 1], lines[i + 2])
            i += 3
        else:
            i += 1
    return refs


def read_references_csv(path: str | Path) -> dict[int, tuple[str, str]]:
    """Parse ``references.csv`` into ``{reference_number: (description, supplier)}``.

    The redistributable form of :func:`read_references`: the same table, as a
    three-column CSV with a header row.
    """
    table = pd.read_csv(path)
    return {
        int(row.reference): (str(row.description), str(row.supplier))
        for row in table.itertuples()
    }


def assign_family(description: str) -> str:
    """Map a free-text sample description to a material family label."""
    for family, pattern in FAMILY_PATTERNS:
        if re.search(pattern, description, re.IGNORECASE):
            return family
    return "OTHER"


# ---------------------------------------------------------------------------
# Workbook reading
# ---------------------------------------------------------------------------
def _read_sheet(worksheet) -> tuple[np.ndarray, np.ndarray]:
    """Extract (wavenumber, intensity) from the first two columns of a sheet."""
    k, y = [], []
    for row in worksheet.iter_rows(min_row=1, max_col=2, values_only=True):
        a, b = row[0], row[1]
        if isinstance(a, (int, float)) and isinstance(b, (int, float)):
            k.append(float(a))
            y.append(float(b))
    k_arr = np.asarray(k, dtype=float)
    y_arr = np.asarray(y, dtype=float)
    order = np.argsort(k_arr)
    return k_arr[order], y_arr[order]


def _build_spectra_set(sheets, refs, min_points: int) -> SpectraSet:
    """Assemble a :class:`SpectraSet` from an iterable of native-grid spectra.

    ``sheets`` yields ``(sheet_name, wavenumber, intensity)`` triples on each
    sheet's own grid; they are linearly interpolated onto the common grid so
    that they can be compared. This is the shared core of the two loaders
    below, so that reading the workbook and reading the redistributable
    extract of Sect. "Data" of the README follow exactly the same code path.
    """
    grid = np.linspace(WAVENUMBER_MIN, WAVENUMBER_MAX, N_GRID_POINTS)

    spectra, rows = [], []
    for sheet_name, k, y in sheets:
        if len(k) < min_points:
            continue

        spectra.append(np.interp(grid, k, y))

        if sheet_name in SPECIAL_SHEETS:
            number, description, supplier = SPECIAL_SHEETS[sheet_name]
        elif sheet_name.isdigit():
            number = int(sheet_name)
            description, supplier = refs.get(number, (f"(unlisted: {sheet_name})", "?"))
        else:
            number = -1
            description, supplier = (f"(unlisted: {sheet_name})", "?")

        rows.append(
            {
                "sheet": sheet_name,
                "reference": number,
                "description": description,
                "description_en": descriptions.translate(description),
                "supplier": supplier,
                "family": assign_family(description),
                "n_points_raw": len(k),
                "wavenumber_min_raw": float(k.min()),
                "wavenumber_max_raw": float(k.max()),
            }
        )

    return SpectraSet(grid, np.vstack(spectra), pd.DataFrame(rows))


def load_workbook_spectra(
    xlsx_path: str | Path,
    references_path: str | Path,
    min_points: int = 1000,
) -> SpectraSet:
    """Read every worksheet of the results workbook into a :class:`SpectraSet`.

    Each worksheet corresponds to one material reference; columns A and B hold
    the wavenumber [cm^-1] and the ATR absorbance respectively.
    """
    refs = read_references(references_path)
    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=False)
    sheets = ((name, *_read_sheet(workbook[name])) for name in workbook.sheetnames)
    return _build_spectra_set(sheets, refs, min_points)


def load_native_spectra(
    npz_path: str | Path,
    references_path: str | Path,
    min_points: int = 1000,
) -> SpectraSet:
    """Read the redistributable extract of the workbook.

    ``npz_path`` holds one ``k_<sheet>`` / ``y_<sheet>`` pair per worksheet on
    the instrument's own grid, in the order given by its ``sheet`` array, and
    ``references_path`` the reference table as CSV. The extract carries exactly
    the two columns of each worksheet that :func:`load_workbook_spectra` reads,
    so both routes produce the same :class:`SpectraSet`.
    """
    refs = read_references_csv(references_path)
    data = np.load(npz_path, allow_pickle=False)
    names = [str(name) for name in data["sheet"]]
    sheets = ((name, data[f"k_{name}"], data[f"y_{name}"]) for name in names)
    return _build_spectra_set(sheets, refs, min_points)


# ---------------------------------------------------------------------------
# Normalisation
# ---------------------------------------------------------------------------
def normalise(intensity: np.ndarray, method: str = "minmax") -> np.ndarray:
    """Normalise spectra row-wise so that topological lifetimes are comparable.

    ATR absorbance depends on the contact pressure between sample and crystal,
    so the absolute scale carries no material information and must be removed
    before persistence is computed (persistence values inherit the units of the
    filtration function).

    Parameters
    ----------
    intensity
        Array of shape ``(n_samples, n_points)``.
    method
        ``"minmax"``  -- rescale each spectrum to ``[0, 1]`` (default);
        ``"max"``     -- divide by the maximum absorbance;
        ``"snv"``     -- standard normal variate (zero mean, unit variance);
        ``"none"``    -- return the input unchanged.
    """
    x = np.atleast_2d(np.asarray(intensity, dtype=float))

    if method == "none":
        out = x
    elif method == "max":
        out = x / np.abs(x).max(axis=1, keepdims=True)
    elif method == "minmax":
        lo = x.min(axis=1, keepdims=True)
        hi = x.max(axis=1, keepdims=True)
        out = (x - lo) / (hi - lo)
    elif method == "snv":
        out = (x - x.mean(axis=1, keepdims=True)) / x.std(axis=1, keepdims=True)
    else:
        raise ValueError(f"unknown normalisation method: {method!r}")

    return out
